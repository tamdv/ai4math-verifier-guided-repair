"""Group 6 presentation deck (16:9, Helvetica, white/blue, Morph transitions) -> slides/Slides_Nhom6_VerifierRepair.pptx

Primitives and design rules reused from Tâm's LinCIR deck builder (same author).

Design rules (from Tâm): 16:9; white + fresh blue palette, modern, restrained; Morph between slides; text >= 20 pt;
Helvetica; titles state the point of the slide; footer = slide number only; reuse the existing diagrams.

Morph: every slide carries the same named shapes ("!!orb", "!!dot", "!!title", "!!num"); PowerPoint matches shapes whose
names start with "!!" across consecutive slides and animates position, size and text between them.

Numbers come from results/raw (via results/figures/*.md); figures from results/figures/.
Run:  python slides/build_slides.py   then   ./slides/to_pdf.sh   (PowerPoint export for checking)
"""
import copy
import json
import statistics
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides" / "Slides_Nhom6_VerifierRepair.pptx"
FIG = ROOT / "results" / "figures"
RES = ROOT / "results"

# palette: white dominant, one fresh blue, navy text, light tint for surfaces
BLUE = RGBColor(0x1E, 0x6F, 0xE8)
BLUE_DARK = RGBColor(0x14, 0x4F, 0xB0)
NAVY = RGBColor(0x0F, 0x1C, 0x3F)
TINT = RGBColor(0xEA, 0xF2, 0xFF)
TINT2 = RGBColor(0xD6, 0xE6, 0xFF)
MUTED = RGBColor(0x5B, 0x6B, 0x82)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xB4, 0x6B, 0x00)
FONT = "Helvetica"
MIN_PT = 20

W, H = Inches(13.333), Inches(7.5)

# orb (big light circle) and dot (small blue circle) positions per slide: x, y, diameter (inches) -> Morph motion
ORB = [(7.6, -1.2, 7.4), (10.6, -2.4, 4.6), (-1.6, 4.4, 4.2), (10.9, 4.6, 4.0), (10.6, -2.4, 4.6), (-2.0, -2.0, 4.0),
       (10.9, 4.8, 3.8), (11.2, -2.2, 4.2), (-1.8, 4.6, 4.0), (10.8, 4.9, 3.8), (11.0, -2.3, 4.4), (-2.0, 4.5, 4.2),
       (10.9, 4.8, 3.9), (11.2, -2.4, 4.4), (-1.9, -2.2, 4.0), (10.8, 4.7, 4.0), (6.9, -0.6, 8.2)]
DOT = [(11.6, 5.7, 0.9), (12.25, 0.55, 0.45), (0.06, 6.6, 0.45), (12.8, 3.6, 0.45), (12.25, 0.55, 0.45), (0.06, 3.4, 0.45),
       (12.8, 4.2, 0.45), (12.3, 0.55, 0.45), (0.06, 6.6, 0.45), (12.8, 3.4, 0.45), (12.3, 0.55, 0.45), (0.06, 6.6, 0.45),
       (12.8, 4.0, 0.45), (12.3, 0.55, 0.45), (0.06, 3.6, 0.45), (12.8, 3.8, 0.45), (6.4, 5.4, 1.0)]

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MORPH_XML = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
    'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
    '<p:transition spd="slow" p14:dur="1200"><p159:morph option="byObject"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')


def _sha(path):
    import hashlib
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


# ---------------------------------------------------------------- primitives


def _font(run, size, bold=False, color=NAVY, italic=False):
    run.font.name = FONT
    run.font.size = Pt(max(size, MIN_PT))
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def text(slide, x, y, w, h, parts, size=22, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name=None, line=1.1):
    """parts: str | list of paragraphs; a paragraph is str or list of (text, {bold, color, italic, size})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = parts if isinstance(parts, list) else [parts]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", False))
    return tb


def bullets(slide, x, y, w, h, items, size=22, gap=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.32)))
        pPr.set("indent", str(-Inches(0.32)))
        bu = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buClr")
        etree.SubElement(bu, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr").set("val", "1E6FE8")
        etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buFont").set("typeface", FONT)  # else "•" falls back to Times
        etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar").set("char", "•")
        runs = it if isinstance(it, list) else [(it, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("bold", False), o.get("color", NAVY))
    return tb


def shape(slide, kind, x, y, w, h, fill, name=None, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def card(slide, x, y, w, h, heading, body, fill=TINT, head_color=BLUE_DARK, size=21):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.55, heading, size=22, bold=True, color=head_color)
    if isinstance(body, list) and body and not isinstance(body[0], tuple):
        bullets(slide, x + 0.25, y + 0.8, w - 0.5, h - 0.95, body, size=size, gap=6)
    else:
        text(slide, x + 0.25, y + 0.8, w - 0.5, h - 0.95, body, size=size)


def stat(slide, x, y, w, value, label, color=BLUE):
    text(slide, x, y, w, 0.95, value, size=48, bold=True, color=color)
    text(slide, x, y + 0.95, w, 0.85, label, size=20, color=MUTED)


def image(slide, path, x, y, w=None, h=None):
    if not Path(path).exists():
        ph = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w or 6, h or 3, TINT)
        text(slide, x + 0.3, y + 0.3, (w or 6) - 0.6, 1.2, f"[Chờ kết quả: {Path(path).name}]", size=20, color=AMBER)
        return ph
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def placeholder(slide, x, y, w, h, msg):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, TINT)
    text(slide, x + 0.3, y + 0.25, w - 0.6, h - 0.5, msg, size=22, color=AMBER)


def table(slide, x, y, w, rows, col_w, size=20, header_fill=BLUE, highlight_row=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(0.5 * nr)).table
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = gt.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (TINT if i == highlight_row else WHITE)
            tf = c.text_frame
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = str(v)
            amber = str(v).startswith("[")
            _font(r, size, bold=(i == 0 or i == highlight_row), color=WHITE if i == 0 else (AMBER if amber else NAVY))
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gt


# ---------------------------------------------------------------- slide frame


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def new(self, title, notes="", title_size=32, cover=False):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = WHITE
        i = self.n
        ox, oy, od = ORB[i % len(ORB)]
        shape(s, MSO_SHAPE.OVAL, ox, oy, od, od, TINT, name="!!orb")
        dx, dy, dd = DOT[i % len(DOT)]
        shape(s, MSO_SHAPE.OVAL, dx, dy, dd, dd, BLUE, name="!!dot")
        if not cover:
            text(s, 0.6, 0.38, 11.4, 1.25, title, size=title_size, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, name="!!title")
            text(s, 11.9, 6.82, 1.0, 0.5, str(i + 1), size=20, color=MUTED, align=PP_ALIGN.RIGHT, name="!!num")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.n += 1
        return s

    def save(self):
        global OUT
        stamp = OUT.with_suffix(".generated.sha256")
        if OUT.exists() and (not stamp.exists() or stamp.read_text().strip() != _sha(OUT)):
            # edited by hand in PowerPoint after the last build: never overwrite it
            OUT = OUT.with_name(OUT.stem + "_auto.pptx")
            print(f"[!] Slides đã được sửa tay -> không ghi đè; bản tự sinh ghi vào {OUT.name}")
        for s in self.prs.slides:  # Morph transition on every slide
            el = s._element
            frag = etree.fromstring(MORPH_XML)
            timing = el.find(f"{{{NS_P}}}timing")
            if timing is not None:
                timing.addprevious(frag)
            else:
                ext = el.find(f"{{{NS_P}}}extLst")
                ext.addprevious(frag) if ext is not None else el.append(frag)
        self.prs.save(OUT)
        if OUT.name == "Slides_Nhom6_VerifierRepair.pptx":
            OUT.with_suffix(".generated.sha256").write_text(_sha(OUT))


# ---------------------------------------------------------------- group 6 content


def arrow(slide, x, y, w):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.32))
    s.fill.solid()
    s.fill.fore_color.rgb = BLUE
    s.line.fill.background()
    s.shadow.inherit = False


def flow(slide, y, steps, w=1.9, gap=0.42, x0=0.6, h=1.25, fills=None):
    """Horizontal pipeline of rounded boxes joined by arrows."""
    x = x0
    for i, step in enumerate(steps):
        fill = (fills or {}).get(i, TINT)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
        text(slide, x + 0.08, y + 0.08, w - 0.16, h - 0.16, step.split("\n"), size=20, bold=True,
             color=BLUE_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            arrow(slide, x + w + 0.06, y + h / 2 - 0.16, gap - 0.12)
        x += w + gap


def build():
    d = Deck()

    # 1 cover
    s = d.new("", notes="Tâm: giới thiệu đề tài, môn học, giảng viên, nhóm. Một câu: hệ thống sửa chứng minh Lean 4 hàng loạt, "
              "chỉ chấp nhận bản sửa được Lean kiểm chứng và giữ nguyên đề, dưới giới hạn chi phí và độ trễ.", cover=True)
    text(s, 0.8, 1.0, 8.4, 0.5, "DỮ LIỆU LỚN, INT4418 · NHÓM 6", size=20, bold=True, color=BLUE, name="!!title")
    text(s, 0.8, 1.7, 8.6, 2.8, [[("Verifier-Guided Repair", {"size": 44, "bold": True})],
                                 [("Sửa chứng minh Lean 4 có kiểm chứng dưới ràng buộc chi phí và độ trễ", {"size": 28, "color": BLUE_DARK})]], line=1.05)
    text(s, 0.8, 4.55, 8.6, 2.1, ["Giảng viên: TS. Nguyễn Kiều Linh",
                                  "Đào Văn Tâm · Trần Quang Đức Dũng · Lâm Thành Trung",
                                  "Nguyễn Xuân Tùng · Khamsing OUTHAIHUENG · Mekdala Nounou"], size=20, color=MUTED)

    # 2 problem
    s = d.new("Hàng nghìn chứng minh Lean 4 lỗi cần được sửa tự động và đáng tin",
              notes="Tâm: vị trí Nhóm 6 trong chuỗi AI4Math, nhận ứng viên lỗi từ Nhóm 5, trả chứng minh đã kiểm chứng cho Nhóm 7/8. "
                    "Thử lại nguyên prompt bỏ qua thông tin quý nhất: thông báo lỗi của Lean.")
    flow(s, 2.0, ["Nhóm 5\nứng viên Lean", "Nhóm 6\nphân loại + sửa", "Lean\nkiểm chứng", "Nhóm 7 / 8\nchứng minh đúng"],
         w=2.55, gap=0.55, fills={1: TINT2})
    card(s, 0.6, 3.75, 12.1, 2.7, "Ba câu hỏi nghiên cứu", [
        [("RQ1 ", {"bold": True}), ("Thông báo lỗi có cấu trúc có tốt hơn thử lại nguyên prompt?", {})],
        [("RQ2 ", {"bold": True}), ("Luật hay LLM, thử bao nhiêu lần, dưới ngân sách token và thời gian?", {})],
        [("RQ3 ", {"bold": True}), ("Mở rộng thế nào theo 100 → 1K → 10K ứng viên và 1/2/4 worker?", {})]])

    # 3 big data
    s = d.new("Nút thắt nằm ở quy mô, lập lịch và chi phí, không ở một phép tính",
              notes="Tâm: bốn con số làm nên bài toán Big Data. Kiểm chứng nguội nạp Mathlib ~6 GB mỗi tệp; số ứng viên 10^4; "
                    "mỗi lần gọi LLM tốn token; log phải phát lại được để kiểm toán.")
    stat(s, 0.8, 2.0, 3.0, "63–107 s", "kiểm chứng nguội 1 tệp, nạp Mathlib")
    stat(s, 3.95, 2.0, 2.9, "10.000", "ứng viên ở mức lớn nhất")
    stat(s, 7.0, 2.0, 2.8, "~1.000", "token mỗi lần LLM sửa")
    stat(s, 10.0, 2.0, 2.9, "48", "run record + log phát lại", color=BLUE_DARK)
    card(s, 0.8, 4.4, 11.9, 2.05, "Mục tiêu hệ thống", [
        "Thông lượng cao, độ trễ đuôi P95 thấp trên vài worker CPU",
        "Chi phí token có trần cứng; kết quả tái lập được từ dữ liệu thô"])

    # 4 data
    s = d.new("Dataset 17.031 lỗi có nhãn lấy từ chính Lean, không viết tay",
              notes="Nounou: 5.000 chứng minh Lean Workbook, MIT kiểm chứng lại với Lean 4.35 → 3.666 seed; 1.110 chứng minh cũ hỏng do "
                    "Mathlib đổi tên = lỗi tự nhiên; 9 toán tử gây lỗi; nhãn từ thông báo Lean thật; chia train/hold-out theo định lý gốc.")
    stat(s, 0.8, 1.95, 2.9, "17.031", "mẫu lỗi, dataset v1.0")
    stat(s, 3.85, 1.95, 2.9, "3.666", "định lý gốc đã kiểm chứng")
    stat(s, 6.9, 1.95, 2.9, "1.110", "lỗi tự nhiên do Mathlib đổi")
    stat(s, 9.95, 1.95, 2.9, "94,4%", "nhãn khớp loại lỗi dự kiến", color=BLUE_DARK)
    card(s, 0.8, 4.35, 5.8, 2.15, "Chống rò rỉ", ["Chia 70/30 theo định lý gốc", "0 định lý ở cả hai phía, 0 bản trùng"])
    card(s, 6.85, 4.35, 5.85, 2.15, "Ba mức quy mô", ["100 ⊂ 1.000 ⊂ 10.000 mẫu", "Hold-out 1K: 284 · 10K: 2.907"], fill=TINT2)

    # 5 pipeline
    s = d.new("Mỗi bản sửa phải qua Lean và kiểm toán trước khi được chấp nhận",
              notes="Khamsing + Trung + Tùng: parser lấy lỗi gốc từ thông báo Lean, taxonomy 7 nhóm, thêm TAC_07; router gửi lỗi đơn giản "
                    "cho luật 0 token, còn lại cho LLM có chẩn đoán; Lean REPL kiểm chứng; kiểm toán chặn sorry/axiom/đổi đề; "
                    "vòng lặp tối đa k lần, dừng khi ứng viên lặp lại.")
    flow(s, 2.05, ["Parser +\ntaxonomy", "Router", "Luật / LLM", "Lean REPL", "Kiểm toán"], w=2.0, gap=0.5,
         fills={4: TINT2})
    card(s, 0.6, 3.75, 3.9, 2.75, "7 nhóm lỗi", ["Cú pháp, import, định danh", "Kiểu, timeout, sorry", "Tactic thất bại, nhóm mới"])
    card(s, 4.75, 3.75, 3.9, 2.75, "Vòng lặp có giới hạn", ["Tối đa k = 1 hoặc 3 lần", "300 s mỗi mẫu", "Dừng nếu ứng viên lặp"])
    card(s, 8.9, 3.75, 3.8, 2.75, "Kiểm toán", ["Lean chấp nhận", "Không sorry / axiom", "Giữ nguyên phát biểu"], fill=TINT2)

    # 6 REPL
    s = d.new("Giữ Lean chạy thường trú: nhanh hơn 3–4 bậc độ lớn",
              notes="Dũng: mỗi worker nạp Mathlib một lần rồi kiểm từng định lý trong môi trường gốc, cô lập; khớp 74/74 với chạy nguội. "
                    "Bài học vận hành: lake sinh tiến trình repl con, huỷ lake khi timeout làm rò repl 100% CPU; sửa bằng process group.")
    stat(s, 0.8, 2.0, 3.8, "63–107 s", "mỗi tệp, chạy lean nguội")
    stat(s, 4.8, 2.0, 3.8, "3 ms", "P50 mỗi định lý với REPL", color=BLUE_DARK)
    stat(s, 8.8, 2.0, 3.8, "23 ms", "P95 mỗi định lý với REPL", color=BLUE_DARK)
    card(s, 0.8, 4.35, 5.8, 2.15, "Đúng như chạy nguội", ["Khớp 74/74 lần kiểm", "Mỗi lần kiểm cô lập trạng thái"])
    card(s, 6.85, 4.35, 5.85, 2.15, "Bài học vận hành", ["Rò tiến trình repl khi timeout", "Sửa: huỷ cả process group"], fill=TINT2)

    # 7 setup
    s = d.new("Cùng dữ liệu, phần cứng và seed, chỉ đổi chiến lược",
              notes="Tâm: Kaggle CPU 4 vCPU / 31 GB; Gemini 3.5 Flash-Lite nhiệt độ 0; baseline bắt buộc = thử lại nguyên prompt đúng 1 lần; "
                    "3 lần lặp, báo trung vị và min–max; trần chi phí $8/tháng, tổng chi thực tế $2,97.")
    card(s, 0.6, 1.9, 3.9, 3.5, "Chiến lược", ["Không sửa · chỉ luật", "Blind là baseline", "LLM + chẩn đoán · lai"])
    card(s, 4.75, 1.9, 3.9, 3.5, "Ma trận", ["k = 0 / 1 / 3", "100 / 1K / 10K mẫu", "1 / 2 / 4 worker · 3 lần"])
    card(s, 8.9, 1.9, 3.9, 3.5, "Môi trường", ["Kaggle 4 vCPU, 31 GB", "Lean 4.35 + Mathlib", "Gemini Flash-Lite"], fill=TINT2)
    text(s, 0.6, 5.7, 12.1, 0.8, [[("Tổng chi phí API cho toàn bộ thực nghiệm: ", {"bold": True, "color": BLUE_DARK}), ("$2,97", {"bold": True})]], size=22)

    # 8 RQ1
    s = d.new("RQ1: thông báo lỗi tăng 22 điểm % và còn rẻ hơn 45%",
              notes="Tùng: cùng 284 mẫu hold-out, k = 1. Không có chẩn đoán, mô hình viết lại cả tệp, 783 token ra/lần nên đắt hơn; "
                    "có chẩn đoán, mô hình sửa cục bộ, 389 token. Lợi ích lớn nhất ở SYN_01 47 → 85%, TMO_05 10 → 69%.")
    stat(s, 0.8, 2.0, 3.6, "31,7%", "baseline thử lại nguyên prompt", color=MUTED)
    stat(s, 4.6, 2.0, 3.6, "53,9%", "LLM + chẩn đoán Lean", color=BLUE_DARK)
    stat(s, 8.4, 2.0, 4.3, "$1,13", "chi phí/1.000 mẫu, baseline $2,05")
    card(s, 0.8, 4.35, 11.9, 2.15, "Vì sao rẻ hơn", [
        "Prompt dài hơn 68% nhưng câu trả lời ngắn bằng một nửa",
        "Token ra đắt gấp 8,3 lần token vào → tổng chi phí giảm"])

    # 9 RQ2 ablation
    s = d.new("RQ2: thử 3 lần đạt 65,1%, gấp đôi baseline",
              notes="Tùng/Tâm: cột xanh là biên dịch, cột cam là thành công đã kiểm toán. k = 3 thêm 11 điểm % so với k = 1; "
                    "82,7% thành công đến ngay lần thử đầu, 12,4% ở lần 2. Hybrid kém LLM thuần ở cùng k, xem slide sau.")
    image(s, FIG / "fig_ablation_slide.png", 0.55, 1.75, w=12.2)

    # 9b RQ2 Pareto
    s = d.new("Chẩn đoán vừa tốt hơn vừa rẻ hơn, nằm trên biên Pareto",
              notes="Tâm: trục x là chi phí LLM cho 1.000 mẫu, trục y là tỷ lệ thành công. Biên đi từ chỉ luật $0 qua hybrid k=1, LLM k=1, hybrid k=3 tới LLM k=3. "
                    "Blind và mô hình 3.6 Flash nằm dưới biên: đắt hơn mà kém hơn.")
    image(s, FIG / "fig_cost_quality_slide.png", 0.6, 1.75, w=7.6)
    card(s, 8.5, 1.9, 4.2, 4.6, "Đọc biểu đồ", ["$0: chỉ luật, 23,9%", "$1,13: LLM k=1, 53,9%", "$2,45: LLM k=3, 65,1%", "Blind: $2,05, 31,7%"], fill=TINT2)

    # 10 rules
    s = d.new("Luật đúng thì gần hoàn hảo, luật sai kéo hybrid xuống",
              notes="Trung: luật được thử trước trên 203/284 mẫu nhưng chỉ đúng 33%; với k = 1 luật sai là hết lượt. "
                    "R0, khôi phục import 97%, R7, heartbeat 100%; R2–R6 0% trên tệp hoàn chỉnh. Cải tiến: chỉ dùng luật chính xác, đo trên train.")
    table(s, 0.6, 1.95, 7.4, [["Luật", "Áp dụng", "Chính xác"],
                              ["R0 khôi phục import", "99", "97%"], ["R7 ngân sách heartbeat", "81", "100%"],
                              ["R1 cân bằng ngoặc", "99", "27%"], ["R3 bảng import", "253", "0%"],
                              ["R2 chèn từ khoá", "126", "0%"]], [3.8, 1.6, 2.0], size=20)
    card(s, 8.3, 1.95, 4.4, 4.55, "Hệ quả", ["Hybrid k=1: 30,3%", "LLM k=1: 53,9%", "Router chỉ nên dùng luật chính xác cao"], fill=TINT2)

    # 11 audit
    s = d.new("Không có kiểm toán, tỷ lệ thành công bị thổi phồng",
              notes="Khamsing/Tâm: trong các mẫu thất bại, 17–26% bản sửa của LLM đổi phát biểu và 10–24% chèn sorry hoặc cấu trúc cấm. "
                    "LLM k = 3: biên dịch 75,0% nhưng thành công thật 65,1%.")
    stat(s, 0.8, 2.0, 3.8, "75,0%", "biên dịch được với LLM k=3", color=MUTED)
    stat(s, 4.8, 2.0, 3.8, "65,1%", "thành công đã kiểm toán", color=BLUE_DARK)
    stat(s, 8.8, 2.0, 3.9, "17–26%", "bản sửa lỗi đã đổi đề")
    card(s, 0.8, 4.35, 11.9, 2.15, "Kiểm toán chặn", ["sorry, admit, axiom, native_decide", "Thay đổi tên, giả thiết hoặc kết luận của định lý"])

    # 12 scale
    s = d.new("RQ3: 2 worker nhanh 1,76×, 4 worker chỉ 1,88× vì nghẽn CPU",
              notes="Dũng: 3 mức × 3 mức worker × 3 lần lặp, spread ≤ 4%. Từ 2 lên 4 worker chỉ thêm 7–11% thông lượng, P95 phục vụ "
                    "tăng gấp đôi, 3,1 → 6,5 s ở 10K: 4 REPL tranh 4 vCPU. Bộ nhớ USS 12,6 → 21,1 GB. Mở rộng tiếp cần nhiều máy.")
    image(s, FIG / "fig_scale_speedup_slide.png", 0.5, 1.8, w=6.3)
    stat(s, 7.0, 2.0, 5.6, "1,52 → 2,86", "mẫu/s ở 10K, từ 1 lên 4 worker")
    stat(s, 7.0, 3.8, 5.6, "3,1 → 6,5 s", "P95 thời gian phục vụ, từ 1 lên 4 worker", color=BLUE_DARK)

    # 13 demo
    s = d.new("Demo: lỗi → phân loại → sửa → Lean kiểm tra → kiểm toán",
              notes="Trung/Tùng/Khamsing: chạy slides/demo_script.md, phát lại log, $0. E000880: luật R0, 0 token. "
                    "E000003: cùng loại lỗi, chỉ LLM, lần 1 lách bằng sorry bị chặn, lần 3 đúng, 4.583 token. E000197: đúng ở lần 2.")
    card(s, 0.6, 1.9, 3.9, 4.6, "E000880 · luật", ["Thiếu import Mathlib", "R0 sửa, Lean OK", "0 token · 73 ms"])
    card(s, 4.75, 1.9, 3.9, 4.6, "E000003 · LLM", ["Lần 1: dùng sorry → bị chặn", "Lần 2: không biên dịch", "Lần 3: đúng · 4.583 token"], fill=TINT2)
    card(s, 8.9, 1.9, 3.9, 4.6, "E000197 · vòng lặp", ["Chứng minh là sorry", "Lần 1 sai, lần 2 đúng", "779 token"])

    # 14 conclusion
    s = d.new("Kết luận: chẩn đoán + kiểm toán + REPL thường trú",
              notes="Tâm: trả lời RQ1–RQ3 bằng số; hạn chế trung thực; hướng phát triển.")
    card(s, 0.6, 1.9, 6.0, 4.6, "Kết luận", [
        "RQ1: 31,7% → 53,9%, rẻ hơn 45%",
        "RQ2: k=3 đạt 65,1%; luật cần chọn lọc",
        "RQ3: REPL nhanh 3–4 bậc; bão hoà CPU ở 2 worker"])
    card(s, 6.85, 1.9, 5.85, 4.6, "Hạn chế → hướng phát triển", [
        "Lỗi tổng hợp → dữ liệu thật Nhóm 5",
        "Một họ LLM → mô hình chuyên Lean",
        "Một máy → worker phân tán"], fill=TINT2)

    # 15 roles
    s = d.new("Phân công bảo vệ: mỗi người một mảng, ai cũng nắm toàn pipeline",
              notes="Tâm: giới thiệu người trả lời từng mảng khi hỏi đáp, contributions.md, slides/defense_qa.md.")
    table(s, 0.6, 1.95, 12.1, [["Thành viên", "Mảng trình bày và trả lời"],
                               ["Đào Văn Tâm", "Câu hỏi nghiên cứu, đánh giá, chi phí"],
                               ["Trần Quang Đức Dũng", "Worker pool, REPL, khả năng mở rộng"],
                               ["Lâm Thành Trung", "Luật sửa, độ chính xác từng luật"],
                               ["Nguyễn Xuân Tùng", "LLM: blind vs chẩn đoán, prompt"],
                               ["Khamsing OUTHAIHUENG", "Parser, taxonomy, kiểm toán"],
                               ["Mekdala Nounou", "Dataset, chia tập, data card"]], [4.4, 7.7], size=20)

    # 16 thanks
    s = d.new("", notes="Cả nhóm: cảm ơn, mời hỏi đáp; Tâm điều phối người trả lời theo bảng phân công.", cover=True)
    text(s, 0.8, 2.2, 10.5, 1.0, "Cảm ơn cô và các bạn", size=44, bold=True, name="!!title")
    text(s, 0.8, 3.4, 10.5, 0.8, "Hỏi đáp", size=28, color=BLUE)
    text(s, 11.9, 6.82, 1.0, 0.5, str(d.n), size=20, color=MUTED, align=PP_ALIGN.RIGHT, name="!!num")

    d.save()
    return d.n


if __name__ == "__main__":
    print(OUT, build(), "slides")
