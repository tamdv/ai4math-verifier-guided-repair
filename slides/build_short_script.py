"""Build slides/KichBan_RutGon_5Phut_Nhom6.docx: the 5-minute talk for Slides_Nhom6_RutGon_5Phut.pptx.

Speech comes from slides/short_talk_content.py (same text as the deck's speaker notes); one section per slide.
Timing uses the same syllable count as build_talk_script.py (160 syllables per minute).
Run:  python slides/build_short_script.py
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_talk_script import ACCENT, GREY, heading, mmss, para, seconds, setup, syllables, table  # noqa: E402
from short_talk_content import QA, SLIDES, SPEAKER  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "slides" / "Slides_Nhom6_RutGon_5Phut.pptx"
OUT = ROOT / "slides" / "KichBan_RutGon_5Phut_Nhom6.docx"


def wide(t, widths) -> None:
    """Pin the table width so Pages and Word both use the full text width."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(int(sum(widths) * 567)), "w:type": "dxa"}),
                       ("w:tblLayout", {"w:type": "fixed"})):
        el = tblPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 567)))


def tbl(doc, rows, widths) -> None:
    wide(table(doc, rows, widths), widths)


def titles() -> list[str]:
    out = []
    for s in Presentation(DECK).slides:
        t = next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.name == "!!title"), "")
        out.append(" ".join(t.split()))
    return out


def main() -> None:
    deck = titles()
    if len(deck) < len(SLIDES):
        sys.exit(f"[LỖI] {DECK.name} có {len(deck)} slide, kịch bản cần {len(SLIDES)}; chạy build_slides_short.py")
    deck[0] = "Verifier-Guided Repair"
    timing, clock = [], 0
    for s in SLIDES:
        sec = seconds(sum(syllables(p) for p in s["speech"]))
        clock += sec
        timing.append((sec, clock))

    doc = Document()
    setup(doc)
    para(doc, "Kịch bản rút gọn 5 phút: Nhóm 6", bold=True, color=ACCENT, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"Verifier-Guided Repair · Môn INT4418 Dữ liệu lớn · Người trình bày: {SPEAKER}",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"Bộ slide: slides/{DECK.name}, {len(SLIDES)} slide trình bày và {len(deck) - len(SLIDES)} slide dự phòng. "
              f"Lời nói khoảng {mmss(clock)} ở tốc độ 160 tiếng mỗi phút; cộng chuyển slide và ngắt nghỉ là khoảng 5 phút. "
              "Thời gian còn lại dành cho câu hỏi; mỗi thành viên mở slide dự phòng của mảng mình khi trả lời. "
              "Lời nói của từng slide cũng nằm trong phần Notes của file .pptx.", italic=True, color=GREY, size=11)
    rows = [["Slide", "Tiêu đề", "Thời lượng", "Mốc"]]
    rows += [[str(i), deck[i - 1], mmss(t), mmss(c)] for i, (t, c) in enumerate(timing, 1)]
    tbl(doc, rows, [1.3, 10.5, 2.4, 2.4])

    heading(doc, "Lời trình bày", 1)
    for i, (s, (sec, c)) in enumerate(zip(SLIDES, timing), 1):
        heading(doc, f"Slide {i}. {s['part']} · {mmss(sec)} · mốc {mmss(c)}", 2)
        for text in s["speech"]:
            para(doc, text)

    heading(doc, "Hỏi đáp: ai trả lời, mở slide nào", 1)
    backup = deck[len(SLIDES):]
    tbl(doc, [["Thành viên", "Mảng trả lời", "Slide dự phòng"]] + [list(r) for r in QA], [4.6, 8.0, 4.0])
    tbl(doc, [["Slide dự phòng", "Nội dung"]] + [[t.split(" · ")[0].replace("Dự phòng ", ""), t.split(" · ", 1)[-1]]
                                                    for t in backup], [3.0, 13.6])
    doc.save(OUT)
    print(f"[OK] {OUT.relative_to(ROOT)}: {len(SLIDES)} slide, lời nói {mmss(clock)}")


if __name__ == "__main__":
    main()
