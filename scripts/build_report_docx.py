"""Builds reports/report.docx from reports/report.md (the source of truth).

Supports the Markdown subset used by the report: YAML front matter (title page),
#/## headings, paragraphs with **bold** / *italic* / `code`, "- " lists, pipe
tables, fenced code blocks, italic caption lines and ![caption](image.png).
Open the .docx in Word, right-click the table of contents -> Update Field, then
File -> Save As PDF.

Usage:
    python3 scripts/build_report_docx.py [--md reports/report.md] [--out reports/report.docx]
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Inserts a Word field (PAGE, TOC …)."""
    run = paragraph.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, instr),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)
    if placeholder:
        paragraph.add_run(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def add_inline(paragraph, text: str, size: float | None = None) -> None:
    """Adds text with **bold**, *italic* and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt((size or 12) - 1.5)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 14), ("Heading 2", 12.5)):
        st = styles[name]
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")


def title_page(doc: Document, meta: dict[str, str]) -> None:
    def line(text: str, size: float, bold: bool = False, space: float = 6) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold

    line("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 13, True)
    line("Khoa Trí tuệ Nhân tạo", 12, space=80)
    line(meta.get("subtitle", ""), 13, space=20)
    line(meta.get("title", ""), 18, True, space=40)
    line("Giảng viên hướng dẫn: TS. Nguyễn Kiều Linh", 12, space=10)
    line(meta.get("author", ""), 12, space=80)
    line(meta.get("date", ""), 12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    r = p.add_run("Mục lục")
    r.bold = True
    r.font.size = Pt(14)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              "(Trong Word: chuột phải → Update Field để tạo mục lục)")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, row[j] if j < len(row) else "", size=9.5)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                set_cell_shading(cell, "E8EEF6")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc: Document, lines: list[str]) -> None:
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ln if ln else " ")
        r.font.name = CODE_FONT
        r.font.size = Pt(8.5)
    doc.add_paragraph()


def add_figure(doc: Document, caption: str, path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(15))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(cap, f"*{caption}*", size=10)


def build(md_path: Path, out_path: Path) -> dict[str, int]:
    text = md_path.read_text(encoding="utf-8")
    meta: dict[str, str] = {}
    if text.startswith("---"):
        _, fm, text = text.split("---", 2)
        for ln in fm.strip().splitlines():
            k, v = ln.split(":", 1)
            meta[k.strip()] = v.strip().strip('"')

    doc = Document()
    setup(doc)
    title_page(doc, meta)
    stats = {"headings": 0, "tables": 0, "figures": 0, "words": 0}
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            add_code(doc, lines[i + 1:j])
            i = j + 1
            continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block
                    if not re.match(r"^\|[\s:|-]+\|$", r.strip())]
            add_table(doc, rows)
            stats["tables"] += 1
            continue
        img = re.match(r"^!\[(.*)\]\((.*)\)\s*$", ln)
        if img:
            add_figure(doc, img.group(1), (md_path.parent / img.group(2)).resolve())
            stats["figures"] += 1
        elif ln.startswith("# "):
            doc.add_heading(ln[2:].strip(), level=1)
            stats["headings"] += 1
        elif ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=2)
            stats["headings"] += 1
        elif ln.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, ln[2:])
        elif re.match(r"^\*[^*].*\*$", ln.strip()) and ln.strip().startswith("*Bảng"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, ln.strip(), size=10)
        elif ln.strip():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, ln.strip())
        stats["words"] += len(ln.split())
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="Build report.docx from report.md")
    parser.add_argument("--md", default="reports/report.md")
    parser.add_argument("--out", default="reports/report.docx")
    args = parser.parse_args()
    stats = build(Path(args.md), Path(args.out))
    print(f"[OK] {args.out}: {stats}")


if __name__ == "__main__":
    main()
